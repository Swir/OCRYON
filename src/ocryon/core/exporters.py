from __future__ import annotations

from collections.abc import Sequence
from contextlib import contextmanager
from functools import lru_cache
from io import BytesIO
import math
import os
from pathlib import Path
import tempfile
from typing import TYPE_CHECKING, Iterator

import fitz
from docx import Document as WordDocument

from ocryon.core.document import DocumentPage

if TYPE_CHECKING:
    from ocryon.core.ocr import OCRWord


_UNICODE_FONT_NAME = "OCRYONUnicode"


def recognized_pages(pages: Sequence[DocumentPage]) -> list[DocumentPage]:
    return [page for page in pages if page.has_ocr]


def _reject_source_overwrite(destination: str | Path, sources: Sequence[Path]) -> None:
    """Block deleting the user's imported original via a PDF export.

    samefile checks filesystem identity, not filename spelling, so aliases,
    hardlinks and different case on Windows cannot bypass this safety rule.
    The late check in _atomic_destination closes the normal export-time window
    where a user or another process could point the output at the source.
    """
    path = Path(destination)
    if not path.exists():
        return
    for source in sources:
        try:
            if path.samefile(source):
                raise ValueError("Export cannot replace an original source document")
        except (FileNotFoundError, NotADirectoryError):
            continue


@contextmanager
def _atomic_destination(
    destination: str | Path, *, source_paths: Sequence[Path] = (),
) -> Iterator[Path]:
    """Never truncate an existing document until a complete replacement exists.

    A temporary sibling has the correct extension for ZIP / PDF serializers;
    closing its descriptor first also avoids Windows sharing violations.
    Source identity is checked before staging and directly before commit.
    """
    path = Path(destination)
    _reject_source_overwrite(path, source_paths)
    fd, filename = tempfile.mkstemp(
        dir=str(path.parent), prefix=f".{path.name}.ocryon-", suffix=path.suffix
    )
    os.close(fd)
    pending = Path(filename)
    try:
        yield pending
        _reject_source_overwrite(path, source_paths)
        # Replacing on the same volume is atomic on NTFS and POSIX filesystems.
        os.replace(pending, path)
    finally:
        pending.unlink(missing_ok=True)


@lru_cache(maxsize=1)
def _unicode_font() -> fitz.Font:
    """Obtain a Unicode font already shipped with PyMuPDF.

    This avoids requiring a separately installed host font on a frozen Windows
    system. 'helv' silently corrupts Polish Unicode during extraction.
    """
    return fitz.Font(fontname="cjk")


def _require_unicode_glyphs(font: fitz.Font, text: str) -> None:
    # Stop export instead of silently substituting '.' / U+FFFD in the text
    # layer; in that case the PDF would appear valid but not be searchable.
    if any(not font.has_glyph(ord(char)) for char in set(text) if not char.isspace()):
        raise ValueError("Searchable PDF font cannot represent OCR characters")


def export_txt(pages: Sequence[DocumentPage], destination: str | Path) -> int:
    source_paths = tuple(page.source_path for page in pages)
    selected = recognized_pages(pages)
    if not selected:
        raise ValueError("There are no recognized pages to export.")

    blocks = []
    for page in selected:
        blocks.append(f"=== {page.label} ===\n{page.ocr_text.rstrip()}")

    with _atomic_destination(
        destination, source_paths=source_paths
    ) as pending:
        pending.write_text("\n\n".join(blocks) + "\n", encoding="utf-8")
    return len(selected)


def export_docx(pages: Sequence[DocumentPage], destination: str | Path) -> int:
    source_paths = tuple(page.source_path for page in pages)
    selected = recognized_pages(pages)
    if not selected:
        raise ValueError("There are no recognized pages to export.")

    document = WordDocument()
    document.add_heading("OCRYON OCR Export", level=0)

    for index, page in enumerate(selected):
        if index:
            document.add_page_break()
        document.add_heading(page.label, level=1)
        document.add_paragraph(page.ocr_text)

    with _atomic_destination(
        destination, source_paths=source_paths
    ) as pending:
        document.save(str(pending))
    return len(selected)


def export_searchable_pdf(
    pages: Sequence[DocumentPage],
    destination: str | Path,
) -> int:
    """Export page images with byte-faithful Unicode invisible OCR text."""
    source_paths = tuple(page.source_path for page in pages)
    selected = recognized_pages(pages)
    if not selected:
        raise ValueError("There are no recognized pages to export.")
    _reject_source_overwrite(destination, source_paths)

    font = _unicode_font()
    document = fitz.open()
    try:
        for source_page in selected:
            # Bound OCR text embedded per PDF page, including the word-layer
            # path, before attempting expensive layout/serialization.
            if len(source_page.ocr_text) > 200_000:
                raise ValueError("OCR text exceeds safe searchable PDF layer limit")
            image = source_page.image
            raster_width, raster_height = image.size
            width, height = float(raster_width), float(raster_height)
            if source_page.page_size_pt is not None:
                try:
                    width, height = map(float, source_page.page_size_pt)
                except (TypeError, ValueError, OverflowError):
                    raise ValueError("Invalid physical PDF page dimensions") from None
                if (
                    not all(map(math.isfinite, (width, height)))
                    or width <= 0 or height <= 0
                    or raster_width <= 0 or raster_height <= 0
                ):
                    raise ValueError("Invalid physical PDF page dimensions")
            scale_x, scale_y = width / raster_width, height / raster_height
            pdf_page = document.new_page(width=width, height=height)

            image_buffer = BytesIO()
            image.save(image_buffer, format="PNG")
            pdf_page.insert_image(pdf_page.rect, stream=image_buffer.getvalue())
            pdf_page.insert_font(fontname=_UNICODE_FONT_NAME, fontbuffer=font.buffer)

            # A partly populated / invalid TSV layout must not silently omit
            # recognized words from the searchable PDF's text layer.
            if _complete_word_coverage(
                pdf_page, source_page.ocr_words, source_page.ocr_text,
                scale_x=scale_x, scale_y=scale_y,
            ):
                inserted_words = _insert_invisible_words(
                    pdf_page, source_page.ocr_words, font,
                    scale_x=scale_x, scale_y=scale_y,
                )
            else:
                inserted_words = 0
            if inserted_words == 0:
                text_rect = fitz.Rect(12, 12, max(13, width - 12), max(13, height - 12))
                _insert_invisible_text(pdf_page, text_rect, source_page.ocr_text, font)

            # Pdf text operators can succeed yet silently wrap / fragment a
            # recognized word (or omit unsupported glyphs). Validate the real
            # searchable layer before saving any page to its final location.
            expected = " ".join(source_page.ocr_text.split())
            actual = " ".join(pdf_page.get_text("text").split())
            if actual != expected:
                raise ValueError("Searchable PDF did not preserve complete OCR text")

        # Subset the 3.5 MB embedded Unicode font, preserving ToUnicode maps.
        # This keeps the release's PDF usable and portable without system fonts.
        document.subset_fonts()
        with _atomic_destination(
        destination, source_paths=source_paths
    ) as pending:
            document.save(str(pending), garbage=4, deflate=True)
    finally:
        document.close()

    return len(selected)


def _complete_word_coverage(
    page: fitz.Page, words: Sequence["OCRWord"], expected_text: str,
    *, scale_x: float = 1.0, scale_y: float = 1.0,
) -> bool:
    """Use spatial text only if every word is backed by valid page geometry.

    The GUI can receive truncated TSV boxes or OCR text from another engine.
    A positive count of *some* inserted words is not proof that the complete
    recognized text was preserved in the PDF. Validate before drawing any
    invisible layer; otherwise fallback produces duplicated fragments.
    """
    if not words or not expected_text.strip():
        return False

    width, height = float(page.rect.width), float(page.rect.height)
    fragments = []
    for word in words:
        try:
            fragment = word.text.strip()
            x, top, box_width, box_height = map(
                float, (word.left, word.top, word.width, word.height)
            )
            x, box_width = x * scale_x, box_width * scale_x
            top, box_height = top * scale_y, box_height * scale_y
        except (AttributeError, TypeError, ValueError, OverflowError):
            return False
        if not fragment or not all(
            map(math.isfinite, (x, top, box_width, box_height))
        ) or (
            x < 0 or top < 0 or box_width <= 0 or box_height <= 0
            or x + box_width > width or top + box_height > height
        ):
            return False
        fragments.append(fragment)

    # The original OCR keeps line breaks; compare equivalent word sequences.
    return " ".join(" ".join(fragments).split()) == " ".join(expected_text.split())


def _insert_invisible_words(
    page: fitz.Page,
    words: Sequence["OCRWord"],
    font: fitz.Font,
    *, scale_x: float = 1.0, scale_y: float = 1.0,
) -> int:
    """Place Unicode searchable words near their source coordinates."""
    inserted = 0
    page_width, page_height = float(page.rect.width), float(page.rect.height)

    for word in words:
        text = word.text.strip()
        if not text:
            continue
        _require_unicode_glyphs(font, text)
        try:
            x, top, width, height = map(
                float, (word.left, word.top, word.width, word.height)
            )
            x, width = x * scale_x, width * scale_x
            top, height = top * scale_y, height * scale_y
        except (TypeError, ValueError, OverflowError):
            continue
        if (
            not all(map(math.isfinite, (x, top, width, height)))
            or x < 0 or top < 0 or width <= 0 or height <= 0
            or x + width > page_width or top + height > page_height
        ):
            continue

        max_font_size = max(1.0, min(72.0, height * 0.85))
        unit_width = font.text_length(text, fontsize=1.0)
        if unit_width > 0:
            font_size = max(1.0, min(max_font_size, (width / unit_width) * 0.95))
        else:
            font_size = max_font_size

        baseline = min(page_height - 0.5, top + max(1.0, height * 0.82))
        page.insert_text(
            fitz.Point(x, baseline),
            text,
            fontsize=font_size,
            fontname=_UNICODE_FONT_NAME,
            render_mode=3,
            overlay=True,
        )
        inserted += 1

    return inserted


def _insert_invisible_text(page: fitz.Page, rect: fitz.Rect, text: str, font: fitz.Font) -> None:
    cleaned = text.strip()
    if not cleaned:
        return
    _require_unicode_glyphs(font, cleaned)

    for font_size in (11, 9, 7, 5, 3, 2, 1):
        remaining = page.insert_textbox(
            rect,
            cleaned,
            fontsize=font_size,
            fontname=_UNICODE_FONT_NAME,
            render_mode=3,
            overlay=True,
        )
        if remaining >= 0:
            return

    # Some valid source images are tiny: even 1 pt cannot contain a full
    # paragraph at once. Split whole words into invisible text layers rather
    # than losing recognized words or refusing an otherwise valid scan.
    # All layers live on the same PDF page and are search-indexed separately.
    # Never split in the middle of a word, which would break literal search.
    words = cleaned.split()
    if len(cleaned) > 200_000:
        raise ValueError("OCR text exceeds safe searchable PDF layer limit")
    position = 0
    layer_count = 0
    while position < len(words):
        if layer_count >= 1024:
            raise ValueError("OCR text exceeds safe searchable PDF layer limit")
        chunk_words = min(128, len(words) - position)
        while chunk_words:
            chunk = " ".join(words[position:position + chunk_words])
            remaining = page.insert_textbox(
                rect,
                chunk,
                fontsize=1,
                fontname=_UNICODE_FONT_NAME,
                render_mode=3,
                overlay=True,
            )
            if remaining >= 0:
                break
            chunk_words //= 2
        if chunk_words == 0:
            # Not even one recognized word fits at 1 pt. Never commit a
            # misleading PDF with a silently incomplete OCR text layer.
            raise ValueError("Full searchable OCR text does not fit the PDF page")
        position += chunk_words
        layer_count += 1
